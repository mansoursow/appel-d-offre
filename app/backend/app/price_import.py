# -*- coding: utf-8 -*-
"""Lecture du classeur "Tableau des MI et PTF" -> historique des prix.

Le classeur tenu par l'assistante comporte une feuille par annee ; chaque
feuille repete, mois par mois, la meme ligne d'entete (N, Date, Objet,
STRUCTURE, Agent, Date delai, Observations, Note, Participants, Attribution,
Prix des participants (TTC), Methode de selection).

Seules les lignes ou la colonne des PRIX est renseignee sont retenues : sans
montant, la ligne n'apprend rien sur ce que pratiquent les concurrents. C'est
exactement la demande metier -- savoir, marche par marche, ce que le cabinet a
propose et ce que les autres ont propose.

Ce module ne fait que lire et normaliser. L'enregistrement en base est dans
price_service.py, la ligne de commande dans tools/import_prix_reference.py.
"""
from __future__ import annotations

import hashlib
import re
import unicodedata

import openpyxl

# Entetes rencontrees dans le classeur (une fois normalisees) -> champ interne.
COLUMNS = {
    "objet": "objet",
    "structure": "structure",
    "agent": "agent",
    "date": "date", "date mi": "date",
    "date delai": "date_limite",
    "observations": "observations",
    "note (nmbre de points)": "notes", "note (nmbre de pts)": "notes",
    "participants": "participants", "noms des participants": "participants",
    "attribution": "attribution",
    "prix des participants (ttc)": "prix", "prix des participants": "prix",
    "methode de selection": "methode",
}

DATE_TEXTE = re.compile(r"\d{1,2}[/.]\d{1,2}[/.]\d{4}")
NUMBER = re.compile(r"\d[\d  .]*\d|\d")


def strip_accents(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    return "".join(c for c in decomposed if not unicodedata.combining(c))


def norm(value) -> str:
    """Minuscules, sans accent ni espace superflu : pour comparer des entetes."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", strip_accents(str(value).replace("\xa0", " "))).strip().lower()


def clean(value):
    """Texte affichable : espaces normalises, cellule vide ramenee a None."""
    if value is None:
        return None
    text = re.sub(r"[ \t\xa0]+", " ", str(value)).strip()
    text = re.sub(r"\n\s*", "\n", text).strip()
    return text or None


def canonical(name: str) -> str:
    """Cle de regroupement d'un concurrent : "ADOC SA" et "adoc sa" = pareil."""
    key = strip_accents(name).upper()
    key = re.sub(r"[^A-Z0-9&]+", " ", key)
    return re.sub(r"\s+", " ", key).strip()


def iso_date(value):
    """Date ISO depuis une cellule Excel (datetime) ou un texte JJ/MM/AAAA."""
    if value is None:
        return None
    text = str(value).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        return m.group(0)
    m = re.search(r"(\d{1,2})[/.](\d{1,2})[/.](\d{4})", text)
    if m:
        return "%s-%02d-%02d" % (m.group(3), int(m.group(2)), int(m.group(1)))
    return None


def parse_amount(text: str):
    """Premier montant d'un texte. Renvoie (montant, suspect).

    "suspect" signale une saisie qu'on refuse d'interpreter, par exemple
    "134B440 000" : un chiffre colle a une lettre est presque surement une
    faute de frappe, et deviner donnerait un montant faux (134 au lieu de
    134 440 000). Dans ce cas on n'affiche que le texte brut, sans montant.
    """
    m = NUMBER.search(text)
    if not m:
        return None, False
    end = m.end()
    if end < len(text) and text[end].isalpha() and text[end].upper() != "F":
        return None, True
    digits = re.sub(r"[  .]", "", m.group(0))
    if not digits.isdigit():
        return None, True
    return int(digits), False


def parse_lines(raw):
    """"ADOC: 7 670 000" ligne par ligne -> {nom: reste de la ligne}."""
    result = {}
    if not raw:
        return result
    for line in str(raw).split("\n"):
        line = re.sub(r"[ \t\xa0]+", " ", line).strip()
        if not line or ":" not in line:
            continue
        name, _, rest = line.partition(":")
        name = name.strip(" -.")
        rest = rest.strip()
        if name and rest:
            result[name] = rest
    return result


def build_marche(year: str, objet: str, prix_brut: str, cell) -> dict:
    attribution = clean(cell("attribution"))
    attribution_key = canonical(attribution) if attribution else ""

    notes_by_key = {canonical(k): v for k, v in parse_lines(clean(cell("notes"))).items()}

    offres = []
    for name, rest in parse_lines(prix_brut).items():
        montant, suspect = parse_amount(rest)
        key = canonical(name)
        note_texte = notes_by_key.get(key)
        note = parse_amount(note_texte)[0] if note_texte else None
        offres.append({
            "nom": name,
            "cle": key,
            "montant": montant,
            "montant_texte": rest,
            "montant_incertain": suspect,
            "note": note,
            "est_adoc": "ADOC" in key,
            # L'attribution est saisie en texte libre : "Groupement Enerteam-CECA"
            # doit reconnaitre l'offre "Groupement ENERTEAM-CECA", d'ou la
            # comparaison par inclusion sur les cles normalisees.
            "est_attributaire": bool(attribution_key) and (
                key == attribution_key or key in attribution_key or attribution_key in key
            ),
        })

    observations = clean(cell("observations"))
    # La colonne "Date" du classeur est souvent incoherente (jour et mois
    # inverses) : on prefere la date de depot ecrite dans les observations.
    depot = DATE_TEXTE.search(observations) if observations else None
    date = iso_date(depot.group(0) if depot else None) \
        or iso_date(cell("date_limite")) \
        or iso_date(cell("date"))

    # Cle stable : reimporter le meme classeur ne cree pas de doublon, et
    # corriger une ligne dans Excel remplace bien l'ancienne version.
    empreinte = hashlib.sha1(("%s|%s" % (year, objet)).encode("utf-8")).hexdigest()[:16]
    return {
        "cle": empreinte,
        "annee": int(year),
        "date": date,
        "objet": objet,
        "structure": clean(cell("structure")),
        "methode": clean(cell("methode")),
        "attributaire": attribution,
        "observations": observations,
        "participants_bruts": clean(cell("participants")),
        "offres": offres,
    }


def read_workbook(source) -> list:
    """Liste des marches chiffres du classeur (chemin ou objet fichier)."""
    wb = openpyxl.load_workbook(source, data_only=True)
    marches = []

    for sheet in wb.worksheets:
        year = sheet.title.strip()
        if not re.fullmatch(r"\d{4}", year):
            continue
        cols = {}

        for row in range(1, sheet.max_row + 1):
            values = [sheet.cell(row, c).value for c in range(1, sheet.max_column + 1)]
            normalised = [norm(v) for v in values]

            # Ligne d'entete (repetee a chaque mois) : on remappe les colonnes,
            # leur position changeant d'une annee a l'autre.
            if "objet" in normalised:
                cols = {}
                for index, label in enumerate(normalised):
                    field = COLUMNS.get(label)
                    if field and field not in cols:
                        cols[field] = index
                continue
            if not cols:
                continue

            def cell(field, _values=values, _cols=cols):
                index = _cols.get(field)
                return _values[index] if index is not None and index < len(_values) else None

            objet = clean(cell("objet"))
            prix_brut = clean(cell("prix"))
            if not objet or not prix_brut:
                continue

            marche = build_marche(year, objet, prix_brut, cell)
            if marche["offres"]:
                marches.append(marche)

    marches.sort(key=lambda m: (m["annee"], m["date"] or "", m["objet"]))
    # Un meme objet peut revenir deux fois dans l'annee (lot 1 / lot 2) : la
    # cle doit rester unique, on suffixe les doublons eventuels.
    vus = {}
    for marche in marches:
        n = vus.get(marche["cle"], 0)
        vus[marche["cle"]] = n + 1
        if n:
            marche["cle"] = "%s-%d" % (marche["cle"], n + 1)
    return marches
